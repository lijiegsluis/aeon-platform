"""Deterministic ingestion for the Data Studio.

Two entry points turn documents into *proposed* figures — never trusted facts:

  * parse_file(filename, data, company)  — an uploaded .xlsx/.xlsm/.csv/.pdf
  * collect_web(company)                 — a live fetch from public sources

Both return a list of proposal dicts, each carrying an fy, a canonical line
item, a value, unit, currency, a human-readable source label, a source_url and a
confidence in [0, 1]. Nothing here is trusted: the caller stages every proposal
and applies the user's confidence threshold before anything reaches the model.

Parsing is fully deterministic (no LLM, no API key) and reuses the
standardisation vocabulary so raw accounting labels map into Aeon Nimbus's template.
Because even the open-web path only ever *proposes* sourced figures, the
never-fabricate contract holds: an unsourced or unparseable value is dropped, not
shown.
"""

from __future__ import annotations

import copy
import csv
import io
import re
from typing import Any, Optional

from aeon_nimbus.standardisation import (
    clean_negative_values,
    detect_currency_and_units,
    standardise_fiscal_year,
    standardise_line_item,
    LINE_ITEM_SYNONYMS,
    _normalise_label,
)

# The line items that belong in the historical financials schema (what the deep
# object, controls and Excel model consume). Proposals for anything else are
# dropped so the studio only ever offers figures the model can actually use.
FINANCIAL_ITEMS = {
    "revenue", "ebitda", "ebit", "net_income", "total_assets", "total_equity",
    "total_debt", "cash", "net_debt", "operating_cash_flow", "capex",
    "free_cash_flow", "dividends_paid",
}
_STATEMENT_OF = {
    "revenue": "income_statement", "ebitda": "income_statement", "ebit": "income_statement",
    "net_income": "income_statement", "total_assets": "balance_sheet", "total_equity": "balance_sheet",
    "total_debt": "balance_sheet", "cash": "balance_sheet", "net_debt": "balance_sheet",
    "operating_cash_flow": "cash_flow", "capex": "cash_flow", "free_cash_flow": "cash_flow",
    "dividends_paid": "cash_flow",
}
_YEAR = re.compile(r"(20\d{2}|19\d{2})")

# Identity / common-header aliases the shared vocabulary does not carry (it maps
# turnover/sales → revenue but not the literal header "Revenue"). Kept local so
# platform-wide standardisation behaviour is untouched.
_ALIASES = {
    "revenue": "revenue", "total revenue": "revenue", "net revenue": "revenue",
    "group revenue": "revenue", "total income": "revenue", "operating income": "ebit",
    "net profit": "net_income", "profit after tax": "net_income",
    "profit for the period": "net_income", "profit attributable to owners": "net_income",
    "total borrowings": "total_debt", "total debt": "total_debt", "net debt": "net_debt",
    "cash and equivalents": "cash", "total cash": "cash", "total assets": "total_assets",
    "total equity": "total_equity", "shareholders equity": "total_equity",
}


def _proposal(fy: str, item: str, value: float, *, source: str, source_url: Optional[str],
              confidence: float, unit: Optional[str], currency: Optional[str]) -> dict[str, Any]:
    return {"fy": fy, "item": item, "statement": _STATEMENT_OF.get(item),
            "value": float(value), "unit": unit, "currency": currency,
            "source": source, "source_url": source_url, "confidence": round(confidence, 2)}


def _match_confidence(raw_label: str, base: float) -> tuple[Optional[str], float]:
    """Map a raw label to a canonical item and score how sure the match is.

    An exact vocabulary hit is worth more than a substring hit; the base differs
    by document type (a page-referenced PDF line is less certain than a clean
    spreadsheet cell). Returns (canonical_item_or_None, confidence)."""
    norm = _normalise_label(raw_label)
    if norm in _ALIASES:
        return _ALIASES[norm], min(0.98, base + 0.10)
    item = standardise_line_item(raw_label)
    if item is None:
        return None, 0.0
    exact = norm in LINE_ITEM_SYNONYMS
    return item, min(0.98, base + (0.10 if exact else 0.0))


def _dedupe(props: list[dict]) -> list[dict]:
    """Keep the highest-confidence proposal per (fy, item)."""
    best: dict[tuple[str, str], dict] = {}
    for p in props:
        key = (p["fy"], p["item"])
        if key not in best or p["confidence"] > best[key]["confidence"]:
            best[key] = p
    return sorted(best.values(), key=lambda p: (p["fy"], p["item"]))


# --- tabular sources (xlsx / csv) -----------------------------------------

def _cell_text(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _parse_table(rows: list[list[Any]], *, source: str, source_url: Optional[str],
                 unit: Optional[str], currency: Optional[str], base_conf: float) -> list[dict]:
    """Parse a rectangular block (list of rows) of a financial statement.

    Finds the header row (the one with the most year tokens), maps each column to
    a fiscal year, then reads every labelled line whose label maps into the
    template. Values run through the same negative/bracket cleaner as the rest of
    the platform."""
    if not rows:
        return []
    # 1) locate the header row: the row carrying the most distinct year tokens.
    header_idx, year_cols = -1, {}
    best_years = 0
    for i, row in enumerate(rows[:25]):
        cols: dict[int, str] = {}
        for j, cell in enumerate(row):
            m = _YEAR.search(_cell_text(cell))
            if m:
                cols[j] = f"FY{m.group(1)}"
        if len(cols) > best_years:
            best_years, header_idx, year_cols = len(cols), i, cols
    if not year_cols:
        return []
    # 2) read line items below the header.
    props: list[dict] = []
    for row in rows[header_idx + 1:]:
        label = next((_cell_text(c) for c in row if _cell_text(c) and not _YEAR.search(_cell_text(c))
                      and any(ch.isalpha() for ch in _cell_text(c))), "")
        if not label:
            continue
        item, conf = _match_confidence(label, base_conf)
        if item is None or item not in FINANCIAL_ITEMS:
            continue
        for col, fy in year_cols.items():
            if col >= len(row):
                continue
            val = clean_negative_values(row[col])
            if val is None:
                continue
            props.append(_proposal(fy, item, val, source=f"{source} · row '{label[:40]}'",
                                   source_url=source_url, confidence=conf, unit=unit, currency=currency))
    return props


def _parse_workbook(data: bytes, filename: str, company: dict) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    unit, currency = company.get("unit"), company.get("currency")
    props: list[dict] = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        text = " ".join(_cell_text(c) for row in rows for c in row)
        det = detect_currency_and_units(text)
        props += _parse_table(rows, source=f"uploaded: {filename} · sheet '{ws.title}'",
                              source_url=None, unit=det["unit"] or unit,
                              currency=det["currency"] or currency, base_conf=0.82)
    wb.close()
    return props


def _parse_csv(data: bytes, filename: str, company: dict) -> list[dict]:
    text = data.decode("utf-8", errors="replace")
    rows = [row for row in csv.reader(io.StringIO(text))]
    det = detect_currency_and_units(text)
    return _parse_table(rows, source=f"uploaded: {filename}", source_url=None,
                        unit=det["unit"] or company.get("unit"),
                        currency=det["currency"] or company.get("currency"), base_conf=0.82)


# --- PDF source ------------------------------------------------------------

def _fallback_years(company: dict, n: int) -> list[str]:
    """Most recent n fiscal years already on file, newest first."""
    fys = sorted((f.get("fy") for f in company.get("financials", []) if f.get("fy")), reverse=True)
    return fys[:n]


def _pdf_year_headers(doc, page: int, n: int = 2) -> list[str]:
    """Read the column-year headers at the top of a statement page, in page order."""
    from aeon_nimbus import extract_statements as X
    seen: list[str] = []
    for line in X._rows_by_y(doc, page)[:14]:
        for m in _YEAR.finditer(line):
            fy = f"FY{m.group(1)}"
            if fy not in seen:
                seen.append(fy)
    return seen[:n]


def _parse_pdf(data: bytes, filename: str, company: dict, source_url: Optional[str] = None) -> list[dict]:
    import fitz
    from aeon_nimbus import extract_statements as X
    doc = fitz.open(stream=data, filetype="pdf")
    unit, currency = company.get("unit"), company.get("currency")
    props: list[dict] = []
    try:
        located = X.locate_statements(doc)
        for stmt, wanted in X.KEY_LINES.items():
            page = located.get(stmt)
            if not page:
                continue
            years = _pdf_year_headers(doc, page) or _fallback_years(company, 2)
            for row in X.extract_line_items(doc, page, wanted):
                item, conf = _match_confidence(row.get("matched") or row.get("raw_label", ""), 0.72)
                if item is None or item not in FINANCIAL_ITEMS:
                    item2, conf2 = _match_confidence(row.get("raw_label", ""), 0.72)
                    if item2 is None or item2 not in FINANCIAL_ITEMS:
                        continue
                    item, conf = item2, conf2
                label = source_url or f"uploaded: {filename}"
                for fy, val in zip(years, row.get("values", [])):
                    if val is None:
                        continue
                    props.append(_proposal(fy, item, val, source=f"{label} · {stmt} p.{row['page']}",
                                           source_url=source_url, confidence=conf, unit=unit, currency=currency))
    finally:
        doc.close()
    return props


# --- public API ------------------------------------------------------------

def parse_file(filename: str, data: bytes, company: dict) -> list[dict]:
    """Parse an uploaded financial file into staged proposals."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        props = _parse_workbook(data, filename, company)
    elif name.endswith(".csv"):
        props = _parse_csv(data, filename, company)
    elif name.endswith(".pdf"):
        props = _parse_pdf(data, filename, company)
    else:
        raise ValueError(f"unsupported file type: {filename!r} (use xlsx, xlsm, csv or pdf)")
    return _dedupe(props)


def collect_web(company: dict, *, fetch=None) -> list[dict]:
    """Best-effort live collection from public sources, newest figures proposed
    with their source URL. The official filing (a PDF at annual_report_url) is the
    reliable path; public aggregators are attempted opportunistically. `fetch` is
    an injectable HTTP getter (url) -> (content_bytes, content_type) for testing."""
    fetch = fetch or _http_get
    props: list[dict] = []
    # 1) the company's own official filing — the highest-integrity web source.
    #    Prefer a dedicated financial-statements PDF over the narrative annual report
    #    (the latter often carries no statements). First filing that parses wins.
    for url in (company.get("financials_url"), company.get("annual_report_url")):
        if not url:
            continue
        try:
            content, ctype = fetch(url)
            if content and ("pdf" in (ctype or "").lower() or url.lower().endswith(".pdf")):
                got = _parse_pdf(content, url.rsplit("/", 1)[-1], company, source_url=url)
                if got:
                    props += got
                    break
        except Exception:
            continue
    # 2) public financial-data aggregators, keyed by exchange + ticker (server-rendered HTML).
    q = _aggregator_query(company)
    if q:
        for tmpl in _AGGREGATOR_URLS:
            u = tmpl.format(q=q)
            try:
                content, ctype = fetch(u)
                if content and "html" in (ctype or "").lower():
                    props += _parse_html_tables(content.decode("utf-8", "replace"), company, u)
            except Exception:
                continue
            if props:  # one aggregator that parses is enough
                break
    return _dedupe([p for p in props if _plausible(p)])


# aggregate line items reported in millions — never a single-digit value; a tiny
# magnitude means a per-share figure (EPS/DPS) was captured by mistake.
_AGGREGATE_ITEMS = {
    "revenue", "ebitda", "ebit", "net_income", "total_assets", "total_equity",
    "total_debt", "cash", "net_debt", "operating_cash_flow", "capex",
    "free_cash_flow", "dividends_paid",
}


def _plausible(p: dict) -> bool:
    v = p.get("value")
    if v is None:
        return False
    if p.get("item") in _AGGREGATE_ITEMS and abs(v) < 20:
        return False  # implausibly small for a millions-scale aggregate → misparse
    return True


_AGGREGATOR_URLS = [
    "https://stockanalysis.com/quote/{q}/financials/",
]

# substring of the universe 'exchange' string -> stockanalysis.com exchange code.
# nase/jse are verified; the rest are best-effort (a wrong code just 404s → no data,
# never a wrong figure).
_SA_EXCHANGE = [
    # Africa. "nse" alone is deliberately NOT mapped here — it collides with
    # India's National Stock Exchange, which also goes by "NSE"; "nairobi"
    # disambiguates the Kenyan one instead.
    ("nairobi", "nase"),
    ("johannesburg", "jse"), ("jse", "jse"),
    ("nigerian", "ngx"), ("ngx", "ngx"),
    ("egyptian", "egx"), ("egx", "egx"),
    ("ghana", "gse"),
    ("casablanca", "cse"),
    ("brvm", "brvm"),
    # Europe
    ("london", "lon"), ("lse", "lon"),
    ("euronext paris", "epa"), ("euronext amsterdam", "ams"),
    ("xetra", "etr"), ("deutsche börse", "etr"), ("frankfurt", "etr"),
    ("borsa italiana", "bit"), ("milan", "bit"),
    ("swiss", "swx"), ("six swiss", "swx"),
    # Asia-Pacific
    ("national stock exchange", "nse"), ("nse india", "nse"),
    ("bombay", "bse"), ("bse", "bse"),
    ("hong kong", "hkg"), ("hkex", "hkg"),
    ("shanghai", "sse"), ("shenzhen", "sze"),
    ("tokyo", "tyo"), ("tse", "tyo"),
    ("korea exchange", "krx"), ("kospi", "krx"),
    ("singapore", "sgx"), ("sgx", "sgx"),
    ("australian securities exchange", "asx"), ("asx", "asx"),
    ("taiwan stock exchange", "twse"),
    ("bursa malaysia", "klse"),
    ("stock exchange of thailand", "set"),
    ("indonesia stock exchange", "idx"),
    ("philippine stock exchange", "pse"),
    # Americas
    ("toronto", "tsx"), ("tsx", "tsx"),
    ("b3", "bvmf"), ("bovespa", "bvmf"), ("são paulo", "bvmf"),
    ("mexican stock exchange", "bmv"), ("bolsa mexicana", "bmv"),
]


def _aggregator_query(company: dict) -> Optional[str]:
    """Build the stockanalysis path segment `<exchange>/<TICKER>` for a company.
    Most non-US tickers need the exchange prefix (e.g. `nase/SCOM`, `lon/VOD`,
    `tsx/SHOP`); a bare ticker only resolves for US-style listings."""
    ticker = (company.get("ticker") or "").split(".")[0].split(";")[0].strip().upper()
    if not ticker:
        return None
    exch = (company.get("exchange") or "").lower()
    for key, code in _SA_EXCHANGE:
        if key in exch:
            return f"{code}/{ticker}"
    return ticker


_BROWSER_UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
               "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36")


def _http_get(url: str) -> tuple[Optional[bytes], Optional[str]]:
    """Fetch a URL with a real browser identity. Several African-exchange sites sit
    behind CDN bot-walls (CloudFront/Cloudflare) that 403 a bare user-agent; on a 403
    we warm cookies by visiting the site root first — mirroring a browser's first hop —
    and retry once with a Referer."""
    import requests
    from urllib.parse import urlsplit
    s = requests.Session()
    s.headers.update({
        "User-Agent": _BROWSER_UA,
        "Accept": "text/html,application/xhtml+xml,application/pdf,*/*",
        "Accept-Language": "en-US,en;q=0.9",
    })
    r = s.get(url, timeout=30)
    if r.status_code == 403:
        p = urlsplit(url)
        try:
            s.get(f"{p.scheme}://{p.netloc}/", timeout=20)
        except Exception:
            pass
        r = s.get(url, timeout=30, headers={"Referer": f"{p.scheme}://{p.netloc}/"})
    r.raise_for_status()
    return r.content, r.headers.get("content-type", "")


def _parse_html_tables(html: str, company: dict, url: str) -> list[dict]:
    """Parse server-rendered financial tables (pandas.read_html; degrades to [])."""
    try:
        import pandas as pd
    except Exception:
        return []
    try:
        tables = pd.read_html(io.StringIO(html))
    except Exception:
        return []
    props: list[dict] = []
    for t in tables:
        rows = [list(t.columns)] + t.values.tolist()
        props += _parse_table(rows, source=f"{url}", source_url=url,
                              unit=company.get("unit"), currency=company.get("currency"), base_conf=0.70)
    return props


# --- merge into the working dataset ---------------------------------------

def merge_facts(extracted: dict, approved: list[dict]) -> dict:
    """Merge approved proposals into a copy of a company's extracted dataset.

    Figures land on the matching fiscal-year object (created if new); net debt is
    re-derived when debt and cash are present but net debt was not itself given.
    Returns a new dict — the caller persists it as the working copy."""
    ext = copy.deepcopy(extracted) if extracted else {}
    fins = ext.setdefault("financials", [])
    by_fy = {f.get("fy"): f for f in fins}
    for a in approved:
        fy = a["fy"]
        f = by_fy.get(fy)
        if f is None:
            f = {"fy": fy}
            fins.append(f)
            by_fy[fy] = f
        f[a["item"]] = a["value"]
        if a.get("source"):
            f["source"] = a["source"]
        if a.get("confidence") is not None:
            f["confidence"] = a["confidence"]
        if a.get("unit") and not ext.get("unit"):
            ext["unit"] = a["unit"]
        if a.get("currency") and not ext.get("currency"):
            ext["currency"] = a["currency"]
    for f in fins:
        if f.get("total_debt") is not None and f.get("cash") is not None and f.get("net_debt") is None:
            f["net_debt"] = f["total_debt"] - f["cash"]
    ext["financials"] = sorted(fins, key=lambda x: x.get("fy", ""))
    return ext

"""Do the dashboard and the workbook say the same thing about the same company?

Two valuation paths exist: the dashboard computes a rating from the extracted
record, and the workbook computes one from the model. A reader who opens both
must not be told two different things.
"""
import warnings
warnings.filterwarnings("ignore")
from pathlib import Path

import formulas
from openpyxl import load_workbook

from aeon_nimbus import db as D, platform_data


def workbook_view(path: Path):
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    wb = load_workbook(path)
    name = path.name.upper()

    def val(cell):
        for k, v in sol.items():
            if k.upper().endswith(f"'[{name}]VALUATION'!{cell.upper()}"):
                try:
                    return v.value[0, 0]
                except Exception:
                    return v
        return None

    vs = wb["Valuation"]
    out = {}
    for r in range(1, vs.max_row + 1):
        a = str(vs.cell(row=r, column=1).value or "").strip().lower()
        if a.startswith("weighted target price"):
            out["target"] = val(f"D{r}")
        elif a.startswith("recommendation"):
            out["stance"] = val(f"D{r}")
        elif a.startswith("upside to target"):
            out["upside"] = val(f"D{r}")
    return out


s = D.SessionLocal()
rows, st_dis, tg_dis = [], 0, 0
for c in sorted(s.query(D.Company).all(), key=lambda x: x.name):
    p = Path("output/models") / f"{c.slug}_model.xlsx"
    if not p.exists():
        continue
    deep = platform_data.deep_from_extracted(
        platform_data.merged_extracted(c.extracted or {}), c.universe or {})
    dash = (deep or {}).get("rating") or {}
    try:
        wbv = workbook_view(p)
    except Exception as e:
        rows.append((c.name, dash.get("stance"), dash.get("target"), f"ERR {type(e).__name__}", None))
        continue
    ws_st = str(wbv.get("stance") or "")
    ws_st = ws_st.split("—")[0].strip() if "—" in ws_st else ws_st
    d_st = dash.get("stance")
    d_tg, w_tg = dash.get("target"), wbv.get("target")
    same_st = (d_st or "").lower() == ws_st.lower()
    same_tg = (d_tg and w_tg and abs(d_tg - w_tg) <= max(0.01, abs(w_tg) * 0.02))
    if d_st and ws_st and not same_st:
        st_dis += 1
    if d_tg and w_tg and not same_tg:
        tg_dis += 1
    rows.append((c.name, d_st, d_tg, ws_st, w_tg))

print(f"{'company':32} {'dash':>9} {'dash tgt':>12} {'workbook':>11} {'wb tgt':>12}  ")
print("-" * 88)
for n, ds, dt, ws, wt in rows:
    flag = ""
    if ds and ws and ds.lower() != ws.lower():
        flag = "  <-- STANCE"
    elif dt and wt and abs(dt - wt) > max(0.01, abs(wt) * 0.02):
        flag = "  <-- TARGET"
    print(f"{n[:32]:32} {str(ds):>9} {(f'{dt:,.2f}' if dt else '-'):>12} "
          f"{str(ws)[:11]:>11} {(f'{wt:,.2f}' if isinstance(wt,(int,float)) else '-'):>12}{flag}")
print(f"\nstance disagreements: {st_dis}    target disagreements: {tg_dis}    of {len(rows)}")

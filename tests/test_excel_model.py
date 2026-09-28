from openpyxl import load_workbook

from aeon_nimbus.excel_model import compile_model
from tests.test_controls import GOOD

UNIVERSE = {"ticker": "TST", "exchange": "NSE", "country": "Kenya", "sector": "Cement",
            "sub_sector": "Cement", "peers": [{"name": "PeerCo", "ticker": "PC"}]}

# consolidated tab set
REQUIRED_SHEETS = ["Cover", "Assumptions", "Financials", "Valuation",
                   "Comps_and_KPIs", "Overview", "Sources_and_Controls"]


def _build(tmp_path):
    out = compile_model(GOOD, UNIVERSE, tmp_path / "m.xlsx")
    return out, load_workbook(out["path"])  # reload proves integrity


def test_workbook_has_required_sheets(tmp_path):
    _out, wb = _build(tmp_path)
    for s in REQUIRED_SHEETS:
        assert s in wb.sheetnames
    assert len(wb.sheetnames) <= 8  # consolidated, not sprawling


def test_named_ranges_present_and_absolute(tmp_path):
    _out, wb = _build(tmp_path)
    for n in ["WACC", "Value_Per_Share", "Rev_Growth", "EBITDA_Margin", "Terminal_Growth_Rate", "Shares_Out", "Scenario"]:
        dn = wb.defined_names.get(n)
        assert dn is not None, f"named range {n} missing"
        assert "$" in dn.value, f"named range {n} must be absolute, got {dn.value}"


def test_forecast_cells_are_formulas_not_static(tmp_path):
    _out, wb = _build(tmp_path)
    fin = wb["Financials"]
    # GOOD has 3 historical years -> forecast Y1 is column F (6); EBITDA is row 8.
    ebitda_fc = fin.cell(row=8, column=6).value
    assert isinstance(ebitda_fc, str) and ebitda_fc.startswith("=") and "EBITDA_Margin" in ebitda_fc
    vps = wb["Valuation"].cell(row=15, column=2).value
    assert isinstance(vps, str) and vps.startswith("=")


def test_historical_hardcodes_are_real_values(tmp_path):
    _out, wb = _build(tmp_path)
    fin = wb["Financials"]
    # revenue row 4, three historical years -> 100,110,121
    assert [fin.cell(row=4, column=3 + j).value for j in range(3)] == [100, 110, 121]


def test_controls_and_dq_written(tmp_path):
    out, wb = _build(tmp_path)
    assert out["export_allowed"] is True
    assert len(out["controls"]) > 0
    sc = wb["Sources_and_Controls"]
    cells = [str(sc.cell(row=r, column=c).value) for r in range(1, 40) for c in range(1, 4)]
    assert any("Data-quality score" in x for x in cells)
    assert any(x in ("critical", "warning", "informational") for x in cells)  # control rows written
